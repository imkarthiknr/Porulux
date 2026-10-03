from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field

from services import ai
from services.ai import Credentials
from services.bank_import import parse_amount

HOLDING_TYPES = ("STOCK", "MF", "ETF", "BOND", "SGB", "OTHER")

# ── Column vocabulary across Zerodha, Groww, Upstox, Angel One, ICICI Direct, CAS and friends ───────
# Headers are compared after normalisation (lowercase, punctuation stripped, spaces collapsed).
SYNONYMS: dict[str, list[str]] = {
    "name": ["stock name", "company name", "scheme name", "security name", "fund name", "instrument name",
             "scrip name", "name", "instrument", "security", "scrip", "stock", "symbol", "stock symbol"],
    "symbol": ["symbol", "stock symbol", "trading symbol", "instrument", "scrip symbol", "ticker", "scrip code"],
    "isin": ["isin", "isin code", "isin no", "isin number"],
    "qty": ["qty", "quantity", "units", "net qty", "net quantity", "total quantity", "holding qty",
            "current qty", "free qty", "quantity available", "balance units", "closing units", "unit balance"],
    "avg": ["avg cost", "average cost", "average price", "average buy price", "avg price", "avg buy price",
            "average cost price", "buy avg", "avg buy", "purchase price", "average nav", "avg nav", "cost price"],
    "price": ["ltp", "closing price", "close price", "current price", "current market price", "market price",
              "last price", "nav", "cmp", "previous closing price", "prev close", "current nav", "closing nav"],
    "value": ["cur val", "current value", "closing value", "value at market price", "market value", "valuation",
              "present value", "current market value", "valuation amount", "market val"],
    "invested": ["invested value", "buy value", "invested amount", "total cost", "cost value", "purchase value",
                 "amount invested", "cost", "total investment", "invested", "cost of acquisition", "total cost value"],
    "hint": ["category", "asset class", "asset type", "type", "product", "instrument type", "sub category"],
}
_SKIP_NAME = re.compile(r"^(total|grand total|sub ?total|net|summary|portfolio|note|disclaimer)\b", re.I)


@dataclass
class ParsedHolding:
    name: str
    symbol: str
    isin: str | None
    holding_type: str
    units: float
    avg_buy_price: float | None
    current_price: float | None

    @property
    def key(self) -> str:
        return (self.isin or self.symbol or self.name).upper()


@dataclass
class ParseResult:
    holdings: list[ParsedHolding] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    source_hint: str | None = None


def norm_header(cell: str) -> str:
    s = re.sub(r"[^\w\s]", " ", str(cell).lower().replace("₹", " "))
    return re.sub(r"\s+", " ", s).strip()


def infer_type(name: str, isin: str | None, hint: str | None = None) -> str:
    n = f"{name} {hint or ''}".lower()
    isin = (isin or "").upper()
    if "sgb" in n or "sovereign gold" in n:
        return "SGB"
    if re.search(r"\betf\b", n) or "bees" in n or "exchange traded" in n:
        return "ETF"
    if isin.startswith("INF") or re.search(r"\b(fund|scheme|idcw|mutual|direct plan|regular plan|growth plan|fof)\b", n):
        return "MF"
    if re.search(r"\b(bond|ncd|debenture|g-?sec|t-?bill|sdl|gilt)\b", n):
        return "BOND"
    return "STOCK"


def make_holding(
    *, name: str | None, symbol: str | None, isin: str | None, units, avg=None, price=None,
    invested=None, value=None, hint: str | None = None, forced_type: str | None = None,
) -> ParsedHolding | None:
    def num(x):
        if isinstance(x, (int, float)) and not isinstance(x, bool):
            return float(x)
        return parse_amount(str(x)) if x not in (None, "") else None

    u = num(units)
    if not u or u <= 0:
        return None
    name = (name or symbol or "").strip()
    symbol = (symbol or "").strip()
    isin = (isin or "").strip().upper() or None
    if isin and not re.fullmatch(r"[A-Z]{2}[A-Z0-9]{9}[0-9]", isin):
        isin = None
    if not name and not isin:
        return None
    if _SKIP_NAME.match(name):
        return None
    avg_v, price_v = num(avg), num(price)
    inv_v, val_v = num(invested), num(value)
    if avg_v is None and inv_v:
        avg_v = inv_v / u
    if price_v is None and val_v:
        price_v = val_v / u
    htype = forced_type if forced_type in HOLDING_TYPES else infer_type(name, isin, hint)
    # holdings.symbol is NOT NULL: funds have no ticker, so fall back to ISIN, then a trimmed name
    sym = symbol if symbol and (htype != "MF" or not isin) else (isin or symbol or name[:24])
    return ParsedHolding(
        name=name or sym, symbol=sym[:40], isin=isin, holding_type=htype, units=round(u, 6),
        avg_buy_price=round(avg_v, 4) if avg_v is not None and avg_v >= 0 else None,
        current_price=round(price_v, 4) if price_v is not None and price_v >= 0 else None,
    )


def merge_duplicates(items: list[ParsedHolding]) -> list[ParsedHolding]:
    """A fund in two folios (or a stock in two demat lines) becomes one holding: units add up and the
    average cost is weighted by units."""
    out: dict[str, ParsedHolding] = {}
    for h in items:
        cur = out.get(h.key)
        if cur is None:
            out[h.key] = h
            continue
        total = cur.units + h.units
        if cur.avg_buy_price is not None and h.avg_buy_price is not None:
            cur.avg_buy_price = round((cur.avg_buy_price * cur.units + h.avg_buy_price * h.units) / total, 4)
        elif h.avg_buy_price is not None and cur.avg_buy_price is None:
            cur.avg_buy_price = h.avg_buy_price
        cur.current_price = h.current_price if h.current_price is not None else cur.current_price
        cur.units = round(total, 6)
    return list(out.values())


# ── Tabular (CSV / XLSX) ──────────────────────────────────────────────────────

def read_tables(content: bytes, filename: str) -> list[list[list[str]]]:
    """Every sheet (XLSX) or the single table (CSV) as rows of strings."""
    if filename.lower().endswith((".xlsx", ".xlsm")) or content[:2] == b"PK":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        tables = []
        for ws in wb.worksheets:
            rows = [["" if c is None else str(c) for c in r] for r in ws.iter_rows(values_only=True)]
            if rows:
                tables.append(rows)
        return tables
    text = content.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        reader = csv.reader(io.StringIO(text), dialect)
    except csv.Error:
        reader = csv.reader(io.StringIO(text))
    return [list(reader)]


def _find_header(rows: list[list[str]]) -> int | None:
    names, qtys = set(SYNONYMS["name"]) | set(SYNONYMS["isin"]), set(SYNONYMS["qty"])
    for i, row in enumerate(rows[:60]):
        h = {norm_header(c) for c in row}
        if h & names and (h & qtys or any(c.startswith("quantity") for c in h)):
            return i
    return None


def _col(headers: list[str], concept: str) -> int | None:
    for syn in SYNONYMS[concept]:
        if syn in headers:
            return headers.index(syn)
    return None


def parse_table(rows: list[list[str]]) -> ParseResult:
    res = ParseResult()
    hi = _find_header(rows)
    if hi is None:
        return res
    headers = [norm_header(c) for c in rows[hi]]
    c_name, c_sym, c_isin = _col(headers, "name"), _col(headers, "symbol"), _col(headers, "isin")
    c_avg, c_price = _col(headers, "avg"), _col(headers, "price")
    c_val, c_inv, c_hint = _col(headers, "value"), _col(headers, "invested"), _col(headers, "hint")

    # Zerodha Console lists free and pledged shares in separate columns; all of them are still yours.
    split_qty = [i for i, h in enumerate(headers) if h in ("quantity available", "quantity pledged margin", "quantity pledged loan")]
    c_qty = None if split_qty else _col(headers, "qty")
    if not split_qty and c_qty is None:
        return res

    def cell(row: list[str], i: int | None) -> str:
        return row[i].strip() if i is not None and i < len(row) and row[i] is not None else ""

    for n, row in enumerate(rows[hi + 1:], start=hi + 2):
        if not any(str(c).strip() for c in row):
            continue
        if split_qty:
            vals = [parse_amount(cell(row, i)) or 0.0 for i in split_qty]
            units = sum(vals)
        else:
            units = parse_amount(cell(row, c_qty))
        h = make_holding(
            name=cell(row, c_name) or cell(row, c_sym), symbol=cell(row, c_sym) or cell(row, c_name),
            isin=cell(row, c_isin), units=units, avg=cell(row, c_avg) or None, price=cell(row, c_price) or None,
            invested=cell(row, c_inv) or None, value=cell(row, c_val) or None, hint=cell(row, c_hint) or None,
        )
        if h:
            res.holdings.append(h)
        elif any(cell(row, i) for i in (c_name, c_sym, c_isin)) and not _SKIP_NAME.match(cell(row, c_name) or cell(row, c_sym)):
            res.warnings.append(f"Row {n} skipped (no quantity): {(cell(row, c_name) or cell(row, c_sym))[:40]}")
    return res


def parse_tabular(content: bytes, filename: str) -> ParseResult:
    out = ParseResult()
    for table in read_tables(content, filename):
        part = parse_table(table)
        out.holdings += part.holdings
        out.warnings += part.warnings
    out.holdings = merge_duplicates(out.holdings)
    return out


# ── PDF / image via the user's AI key ─────────────────────────────────────────

HOLDINGS_PROMPT = (
    "Extract EVERY investment holding from this Indian brokerage / demat / mutual fund statement "
    "(e.g. NSDL or CDSL CAS, CAMS or KFintech statement, Zerodha, Groww, Upstox, Angel One, ICICI Direct). "
    "Return ONLY a JSON object (no markdown, no explanation):\n"
    '{"source":"<broker or registrar name, or null>","holdings":[{"name":"<security or scheme name>",'
    '"symbol":"<exchange ticker or null>","isin":"<ISIN or null>","units":<number>,'
    '"average_cost":<cost per unit or null>,"invested_value":<total cost or null>,'
    '"current_price":<LTP or NAV per unit or null>,"current_value":<market value or null>,'
    '"type":"STOCK|MF|ETF|BOND|SGB|OTHER"}]}\n'
    "Use plain numbers in INR. Include stocks, mutual fund schemes (one entry per scheme, summing folios of "
    "the same scheme), ETFs, bonds and sovereign gold bonds. Skip totals, cash balances and transaction history."
)


def normalize_ai_result(raw) -> ParseResult:
    if isinstance(raw, list):
        raw = {"holdings": raw}
    if not isinstance(raw, dict) or not isinstance(raw.get("holdings"), list):
        raise ValueError("expected an object with a holdings array")
    res = ParseResult(source_hint=(raw.get("source") or None))
    for item in raw["holdings"]:
        if not isinstance(item, dict):
            continue
        h = make_holding(
            name=item.get("name"), symbol=item.get("symbol"), isin=item.get("isin"), units=item.get("units"),
            avg=item.get("average_cost"), price=item.get("current_price"),
            invested=item.get("invested_value"), value=item.get("current_value"),
            forced_type=str(item.get("type") or "").upper() or None,
        )
        if h:
            res.holdings.append(h)
    res.holdings = merge_duplicates(res.holdings)
    return res


async def extract_holdings(data: bytes, media_type: str, *, creds: Credentials) -> ParseResult:
    raw = await ai.generate(data, media_type, HOLDINGS_PROMPT, creds=creds, max_tokens=32000)
    return normalize_ai_result(ai.parse_json(raw))


# ── Compare with what the user already has ────────────────────────────────────

def _same(a, b, tol) -> bool:
    if a is None or b is None:
        return a is None and b is None
    return abs(float(a) - float(b)) <= tol


def reconcile(parsed: list[ParsedHolding], existing: list[dict], source: str) -> tuple[list[dict], list[dict]]:
    """Mark each parsed holding as new / update / unchanged against the user's current holdings, and list
    holdings from this source that the statement no longer contains. Hand-added holdings (no source) with a
    matching ISIN or symbol are adopted rather than duplicated."""
    pool = [e for e in existing if e.get("source") in (source, None)]
    by_isin = {str(e["isin"]).upper(): e for e in pool if e.get("isin")}
    by_symbol = {str(e["symbol"]).upper(): e for e in pool if e.get("symbol")}
    matched: set[str] = set()

    entries = []
    for h in parsed:
        ex = (by_isin.get(h.isin) if h.isin else None) or by_symbol.get(h.symbol.upper())
        if ex is not None and ex["id"] in matched:
            ex = None
        action, changes = "new", {}
        if ex is not None:
            matched.add(ex["id"])
            for fld, tol in (("units", 1e-5), ("avg_buy_price", 0.005), ("current_price", 0.005)):
                new = getattr(h, fld)
                if new is not None and not _same(ex.get(fld), new, tol):
                    changes[fld] = {"from": ex.get(fld), "to": new}
            adopted = ex.get("source") is None
            action = "update" if changes or adopted else "unchanged"
        invested = round(h.units * h.avg_buy_price, 2) if h.avg_buy_price is not None else None
        value = round(h.units * h.current_price, 2) if h.current_price is not None else None
        entries.append({
            "key": h.key, "name": h.name, "symbol": h.symbol, "isin": h.isin, "holding_type": h.holding_type,
            "units": h.units, "avg_buy_price": h.avg_buy_price, "current_price": h.current_price,
            "invested": invested, "current_value": value, "action": action,
            "existing_id": ex["id"] if ex is not None else None, "changes": changes,
        })

    missing = [
        {"id": e["id"], "name": e.get("name"), "symbol": e.get("symbol"), "units": e.get("units")}
        for e in existing if e.get("source") == source and e["id"] not in matched
    ]
    return entries, missing


# Guess the broker from a file name so the user rarely has to type it.
_FILE_HINTS = [
    ("zerodha", "Zerodha"), ("kite", "Zerodha"), ("console", "Zerodha"), ("groww", "Groww"), ("upstox", "Upstox"),
    ("angel", "Angel One"), ("icici", "ICICI Direct"), ("hdfc", "HDFC Securities"), ("kotak", "Kotak Securities"),
    ("paytm", "Paytm Money"), ("5paisa", "5paisa"), ("nsdl", "NSDL CAS"), ("cdsl", "CDSL CAS"),
    ("cams", "CAMS"), ("kfin", "KFintech"), ("karvy", "KFintech"),
]


def guess_source(filename: str | None) -> str | None:
    n = (filename or "").lower()
    return next((label for key, label in _FILE_HINTS if key in n), None)
